import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import zipfile
import os

def create_title_page_docx():
    doc = docx.Document()
    
    # Page Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    # Title
    title_p = doc.add_paragraph()
    title_run = title_p.add_run("When Confidence Proxies Confound Reasoning Complexity: Pitfalls of Uncertainty-Weighted Credit Assignment in Language Model Reinforcement Learning")
    title_run.font.size = Pt(16)
    title_run.font.bold = True
    title_run.font.name = "Calibri"
    title_p.paragraph_format.space_after = Pt(18)
    
    # Author
    author_p = doc.add_paragraph()
    author_run = author_p.add_run("Author:")
    author_run.font.bold = True
    author_run.font.size = Pt(12)
    doc.add_paragraph("Sham Thakare (Independent Researcher)")
    
    # Correspondence Address & Email
    corr_p = doc.add_paragraph()
    corr_run = corr_p.add_run("Correspondence & Contact Information:")
    corr_run.font.bold = True
    corr_run.font.size = Pt(12)
    doc.add_paragraph("Sham Thakare\nIndependent Researcher\nEmail: shamthakare3000@gmail.com\nWebsite / Portfolio: https://github.com/shamddd/ear_grpo_reasoning")
    
    # Acknowledgments & Disclosure
    ack_p = doc.add_paragraph()
    ack_run = ack_p.add_run("Acknowledgments & IEEE AI Disclosure:")
    ack_run.font.bold = True
    ack_run.font.size = Pt(12)
    
    ack_text = (
        "In accordance with IEEE Author Guidelines and IEEE Publication Services and Products Board (PSPB) Operations Manual policies regarding generative AI systems: "
        "Generative AI assistance (Google Antigravity / Gemini models) was utilized during the development, refactoring, and execution monitoring of experimental training scripts, LaTeX formatting, and statistical analysis pipelines. "
        "All experimental procedures, neural network weight executions, forward/backward passes, statistical hypothesis tests, and evaluation metrics were executed on dedicated hardware, inspected, and verified by the author. "
        "All bibliographic references were verified against authoritative publisher records. The author assumes full responsibility for the integrity, validity, and accuracy of all scientific claims, data, and conclusions presented in this manuscript.\n\n"
        "Conflict of Interest: The author declares no competing financial or non-financial interests.\n"
        "Funding: This research received no external grant funding."
    )
    doc.add_paragraph(ack_text)
    
    out_docx = "submission/ieee_tai/Title_Page.docx"
    doc.save(out_docx)
    print(f"Saved {out_docx}")

def create_anonymized_bundle_zip():
    zip_path = "submission/ieee_tai/Anonymized_Manuscript_LaTeX_Bundle.zip"
    with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as zipf:
        zipf.write("submission/ieee_tai/anonymized_main.tex", arcname="main.tex")
        zipf.write("submission/ieee_tai/references.bib", arcname="references.bib")
        zipf.write("submission/ieee_tai/IEEEtai.cls", arcname="IEEEtai.cls")
    print(f"Saved {zip_path}")

if __name__ == "__main__":
    create_title_page_docx()
    create_anonymized_bundle_zip()
