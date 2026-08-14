import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_anonymized_docx():
    doc = docx.Document()
    
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    # Title
    t_p = doc.add_paragraph()
    t_run = t_p.add_run("When Confidence Proxies Confound Reasoning Complexity: Pitfalls of Uncertainty-Weighted Credit Assignment in Language Model Reinforcement Learning")
    t_run.font.size = Pt(16)
    t_run.font.bold = True
    t_run.font.name = "Calibri"
    t_p.paragraph_format.space_after = Pt(12)
    
    # Author (Anonymized)
    a_p = doc.add_paragraph()
    a_run = a_p.add_run("Anonymous Author(s)")
    a_run.font.italic = True
    a_run.font.size = Pt(11)
    a_p.paragraph_format.space_after = Pt(14)
    
    # Abstract
    abs_h = doc.add_paragraph()
    abs_h_run = abs_h.add_run("Abstract")
    abs_h_run.font.bold = True
    abs_h_run.font.size = Pt(12)
    
    abs_p = doc.add_paragraph(
        "Reinforcement learning from rule-based verifiers (RLVR), notably Group Relative Policy Optimization (GRPO), "
        "has emerged as a cornerstone for post-training Large Language Models (LLMs) on complex reasoning tasks. "
        "A natural hypothesis is that regularizing policy gradient advantages by trajectory-level uncertainty could prevent "
        "policy collapse and filter noisy exploration traces. In this work, we conduct a systematic empirical and architectural "
        "audit of uncertainty-regularized GRPO. First, we show that Monte Carlo (MC) dropout probing becomes degenerate when the "
        "target architecture lacks active dropout in its evaluated forward graph, as confirmed on Qwen2.5. Second, through diagnostic "
        "evaluations on untouched GSM8K and SVAMP datasets (N=100), we find that internal confidence proxies—token predictive entropy "
        "(r = +0.486), mean token negative log-likelihood (r = +0.432), and logit margins (r = +0.495)—are strongly correlated with "
        "derivation length and arithmetic step complexity. In stress tests, internal token entropy misidentifies correct multi-step "
        "reasoning traces as more uncertain than short incorrect errors in 42.1% of paired comparisons. While external consensus disagreement "
        "(Self-Consistency) provides unconfounded offline error discrimination (AUROC = 0.812), a preregistered 5-way controlled RL experiment "
        "demonstrates that using consensus weights for online advantage scaling yields no performance improvement over standard outcome-supervised "
        "GRPO (80.00% vs 80.00%; Delta = 0.00%). Our findings demonstrate that high offline error-predictive validity does not guarantee online "
        "reinforcement learning credit utility, establishing the necessity of negative controls when evaluating confidence-guided post-training algorithms."
    )
    abs_p.paragraph_format.space_after = Pt(12)
    
    # Impact Statement
    imp_h = doc.add_paragraph()
    imp_h_run = imp_h.add_run("Impact Statement")
    imp_h_run.font.bold = True
    imp_h_run.font.size = Pt(12)
    
    imp_p = doc.add_paragraph(
        "Reinforcement learning algorithms for Large Language Models increasingly govern automated reasoning systems in mathematics, coding, "
        "and decision support. Integrating confidence or uncertainty weights into policy gradient updates is frequently proposed to prevent "
        "reward hacking and model drift. This work provides a rigorous methodological warning: internal uncertainty signals frequently track "
        "legitimate reasoning complexity rather than error, causing naive uncertainty-weighted algorithms to suppress valid multi-step exploration. "
        "By establishing that offline error predictors do not automatically improve online policy learning, our study provides a reproducible "
        "diagnostic framework and negative-control protocol to prevent the adoption of unverified or counterproductive credit assignment mechanisms."
    )
    imp_p.paragraph_format.space_after = Pt(14)
    
    # Section 1: Introduction
    s1 = doc.add_paragraph()
    s1.add_run("1. Introduction").font.bold = True
    doc.add_paragraph(
        "Post-training autoregressive Large Language Models (LLMs) using reinforcement learning from verifiable rewards (RLVR) has achieved "
        "substantial milestones in mathematical derivation and formal deduction. Group Relative Policy Optimization (GRPO) estimates baseline-free "
        "advantages by normalizing outcome rewards across a group of sampled rollouts:\n"
        "A_i = (R(y_i) - mean(R)) / (std(R) + eps)\n\n"
        "Because policy gradient methods can suffer from policy collapse and reward sparsity, numerous approaches hypothesize that weighting "
        "advantages by trajectory-level uncertainty U(y_i) can protect the policy from destabilizing updates.\n\n"
        "In this paper, we report a rigorous empirical and architectural investigation into uncertainty-weighted credit assignment for GRPO. We evaluate:\n"
        "1. Architectural Validity (C1): Does MC-dropout probing produce meaningful stochastic representations on zero-dropout architectures?\n"
        "2. Diagnostic Validity (C2): Do internal token-level confidence proxies distinguish mathematical errors from legitimate reasoning complexity?\n"
        "3. Algorithmic Validity (C3): Does a validated offline error predictor (Self-Consistency) improve online GRPO policy optimization over outcome-supervised baselines and negative controls?"
    )
    
    # Section 2: Architectural Audit
    s2 = doc.add_paragraph()
    s2.add_run("2. Estimator Validity on Zero-Dropout Architectures").font.bold = True
    doc.add_paragraph(
        "We audited the internal layer structure of open-weight causal language models, specifically inspecting Qwen/Qwen2.5-0.5B-Instruct (Qwen2ForCausalLM). "
        "We find that the model contains exactly 0 active nn.Dropout modules in its attention and MLP blocks (attention_dropout = 0.0). Consequently, "
        "executing Monte Carlo dropout forward passes with mc_dropout=True produces mathematically deterministic passes (Var(log P) = 0.0000000000, Delta Logit = 0.0).\n\n"
        "When normalized in advantage scaling equations with stability constant eps = 10^-8, floating-point order-of-operations noise (approx 10^-12) "
        "results in a multiplier of exp(-gamma * 10^-4) approx 0.999965, producing policy updates collinear to standard GRPO (cos(Delta theta) = 1.000000). "
        "We conclude that MC-dropout uncertainty estimation becomes degenerate when the target architecture contains no active dropout in the evaluated forward computation graph."
    )
    
    # Section 3: Diagnostic Proxy Benchmark
    s3 = doc.add_paragraph()
    s3.add_run("3. Confidence Proxies and Reasoning Complexity").font.bold = True
    doc.add_paragraph(
        "Evaluating candidate uncertainty proxies across untouched GSM8K confirmatory data (N = 100 independent prompt clusters, 98 degrees of freedom) "
        "reveals a pattern consistent with a reasoning-complexity confound:\n"
        "- Token Predictive Entropy correlates strongly with sequence length (r = +0.486, 95% CI [+0.318, +0.627]) and equation count (r = +0.421).\n"
        "- After controlling for completion length via partial correlation, the association between token entropy and correctness collapses from r = -0.214 to partial r = -0.092 (p = 0.365).\n"
        "- In the Correct-but-Complex Stress Test, token entropy misidentifies correct multi-step reasoning traces as more uncertain than short incorrect errors in 42.1% of paired comparisons.\n"
        "- Self-Consistency (U_SC = 1 - modal_count / K) remains robust to length bias (r = +0.114, partial r = -0.569, p = 8.1 x 10^-10, AUROC = 0.812)."
    )
    
    # Section 4: RL Experiment
    s4 = doc.add_paragraph()
    s4.add_run("4. Preregistered Consistency-Aware RL Experiment").font.bold = True
    doc.add_paragraph(
        "To test whether validated offline self-consistency translates to an effective online policy gradient weight, we preregistered Consistency-Aware GRPO (CA-GRPO):\n"
        "A_tilde_i = A_hat_i * (1.0 + lambda * (c_i - c_bar))\n"
        "where c_i = 1 if extracted answer equals modal answer else 0.\n\n"
        "We benchmarked CA-GRPO against four matched controls across three independent training seeds on held-out test data evaluated under the canonical 256-token generation budget:\n"
        "- Standard-GRPO (G=4): 80.00 +/- 0.00% Pass@1\n"
        "- Compute-Matched-GRPO (G=8): 78.33 +/- 2.89% Pass@1\n"
        "- Random-Weight-Control (G=4): 75.00 +/- 5.00% Pass@1\n"
        "- Permuted-Consistency-Control (G=4): 80.00 +/- 0.00% Pass@1\n"
        "- CA-GRPO (Proposed, G=4): 80.00 +/- 0.00% Pass@1\n\n"
        "CA-GRPO did not improve performance over Standard GRPO in our tested configuration (Delta = 0.00%, Cohen's d = 0.000). Negative controls (permuted and random weighting) "
        "demonstrated that trajectory-specific advantage scaling provided no causal learning advantage over standard outcome-supervised GRPO."
    )
    
    # Limitations & Conclusion
    s5 = doc.add_paragraph()
    s5.add_run("5. Limitations & Conclusion").font.bold = True
    doc.add_paragraph(
        "This investigation is subject to several limitations: (1) RL training evaluations were conducted with N=3 independent training seeds; (2) evaluations focused on "
        "mathematical reasoning domains (GSM8K, SVAMP) using Qwen2.5-0.5B-Instruct; (3) self-consistency requires additional inference compute (K=4 rollouts per prompt); and (4) "
        "our findings evaluate specific advantage-weighting formulations and do not exclude the possibility of alternative consensus-based post-training designs.\n\n"
        "Conclusion: Our findings demonstrate that internal confidence proxies are consistent with a reasoning-complexity confound, and that high offline error discrimination "
        "does not imply online reinforcement learning credit utility. Rigorous negative controls and architectural audits are indispensable for sound RLVR post-training."
    )
    
    out_docx = "submission/ieee_tai/Anonymized_Main_Document.docx"
    doc.save(out_docx)
    print(f"Saved {out_docx}")

if __name__ == "__main__":
    create_anonymized_docx()
